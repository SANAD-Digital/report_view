"""Excel export with openpyxl (docs/architecture.md §4.3, §9.2).

Values are raw (never HTML from a formatter). Numbers keep their numeric type
with a number format built from the system precision; dates are real dates
with the system date format.

Rows follow the screen (``core/engine.py``): every row is written with its outline
level; the rows under a folded group or parent are hidden and the folded row is
marked collapsed, so Excel opens them with its own outline buttons.
"""

import datetime
import re
from io import BytesIO
from typing import Any

import frappe
from frappe import _
from frappe.utils.jinja_globals import is_rtl
from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.cell_range import CellRange

from report_view.core.export.model import column_precision

NUMERIC_TYPES = ("Currency", "Float", "Int", "Percent", "Duration")
_SHEET_TITLE_BAD = re.compile(r"[\\/*?:\[\]]")
MIXED_MARK = "—"
# Text Excel would read as a formula (OWASP CSV injection): written with a quote prefix.
FORMULA_START = ("=", "+", "-", "@", "\t", "\r")


def _number_format(precision: int) -> str:
	return "#,##0" + ("." + "0" * precision if precision > 0 else "")


def _date_format() -> str:
	fmt = frappe.db.get_single_value("System Settings", "date_format") or "yyyy-mm-dd"
	return fmt.lower()


def _cell_value(value: Any, column: dict) -> Any:
	if value is None or value == "":
		return None
	if isinstance(value, dict) and value.get("mixed"):
		return MIXED_MARK
	ft = column["fieldtype"]
	if ft in NUMERIC_TYPES and not isinstance(value, bool):
		try:
			return float(value) if ft != "Int" else int(float(value))
		except TypeError, ValueError:
			return str(value)
	if ft == "Date" and isinstance(value, str) and len(value) >= 10:
		try:
			return datetime.date.fromisoformat(value[:10])
		except ValueError:
			return value
	if ft == "Datetime" and isinstance(value, str):
		try:
			return datetime.datetime.fromisoformat(value)
		except ValueError:
			return value
	if isinstance(value, list):
		return ", ".join(str(v) for v in value)
	if isinstance(value, int | float | bool | datetime.date):
		return value
	return str(value)


def _filters_text(filters: dict) -> str:
	parts = []
	for key, value in filters.items():
		if value in (None, "", []) or (value == 0 and isinstance(value, int) and not isinstance(value, bool)):
			continue
		if isinstance(value, list):
			value = ", ".join(str(v) for v in value)
		parts.append(f"{_(frappe.unscrub(key))}: {value}")
	return " | ".join(parts)


def _group_label(out: dict, by_key: dict[str, dict]) -> str:
	"""Text that names a group row or the group of a subtotal row."""
	if out["label"] is not None:
		return out["label"]
	value = out["group_value"]
	if value is None or value == "":
		return _("Not Set")
	fieldtype = (by_key.get(out["group_key"]) or {}).get("fieldtype")
	if fieldtype == "Date":
		return frappe.utils.formatdate(value)
	if fieldtype == "Datetime":
		return frappe.utils.format_datetime(value)
	if isinstance(value, list):
		return ", ".join(str(v) for v in value)
	return str(value)


def _bands(columns: list[dict], balance_pairs: list[dict]) -> dict[int, tuple[int, str]]:
	"""Column-index -> (span, label) for adjacent debit/credit pairs."""
	index = {c["key"]: i for i, c in enumerate(columns)}
	bands = {}
	for pair in balance_pairs:
		d, c = index.get(pair.get("debit_column")), index.get(pair.get("credit_column"))
		if d is not None and c is not None and abs(d - c) == 1:
			bands[min(d, c)] = (2, pair.get("label") or "")
	return bands


def _cell(ws, value: Any, font=None, fill=None, border=None, alignment=None, number_format=None):
	"""A styled cell of the streamed sheet. Text that Excel would read as a formula
	(``= + - @``, a tab or a carriage return first: OWASP's CSV-injection list) is written
	as text with Excel's quote prefix, never as a formula."""
	cell = WriteOnlyCell(ws, value=value)
	if isinstance(value, str) and value.startswith(FORMULA_START):
		cell.data_type = "s"
		cell.quotePrefix = True
	if font is not None:
		cell.font = font
	if fill is not None:
		cell.fill = fill
	if border is not None:
		cell.border = border
	if alignment is not None:
		cell.alignment = alignment
	if number_format is not None:
		cell.number_format = number_format
	return cell


def build(
	report_name: str,
	nr: dict,
	result: dict,
	model: dict,
	settings: dict,
) -> bytes:
	"""Build the workbook and return its bytes.

	The sheet is streamed (``write_only``): memory stays flat whatever the row count, so
	rows go out in order and the sheet's settings (widths, frozen header, outline, merges)
	are set before the first row."""
	columns = result["columns"]
	layout = model["layout"]
	filters = model["filters"]
	rtl = is_rtl()
	wb = Workbook(write_only=True)
	ws = wb.create_sheet((_SHEET_TITLE_BAD.sub(" ", _(report_name)) or "Report")[:31])
	ws.sheet_view.rightToLeft = rtl
	ws.sheet_properties.outlinePr.summaryBelow = False

	bold = Font(bold=True)
	header_fill = PatternFill("solid", fgColor="F2F2F2")
	thin = Side(style="thin", color="BFBFBF")
	header_border = Border(top=thin, bottom=thin, left=thin, right=thin)
	ncols = max(len(columns), 1)
	number_formats = {
		c["key"]: _number_format(column_precision(c)) for c in columns if c["fieldtype"] in NUMERIC_TYPES
	}
	date_fmt = _date_format()

	# The rows above the column headers: company, title, filters, then an empty row.
	top: list[tuple[Any, Font | None]] = []
	company = filters.get("company") or frappe.defaults.get_user_default("Company")
	if company:
		top.append((str(company), Font(bold=True, size=14)))
	top.append((layout.get("title") or _(report_name), Font(bold=True, size=12)))
	if layout.get("include_filters"):
		text = _filters_text(filters)
		if text:
			top.append((text, None))
	bands = _bands(columns, settings.get("balance_pairs") or [])
	covered: set[int] = set()
	for start, (span, _label) in bands.items():
		covered.update(range(start, start + span))
	band_row = len(top) + 2
	header_row = band_row + 1 if bands else band_row
	first_row = header_row + 1

	for i, col in enumerate(columns):
		width = col.get("width") or 120
		ws.column_dimensions[get_column_letter(i + 1)].width = max(8, min(80, width / 7))
	ws.freeze_panes = f"A{first_row}"
	for row, _top in enumerate(top, 1):
		ws.merged_cells.add(CellRange(min_row=row, min_col=1, max_row=row, max_col=ncols).coord)
	for start, (span, _label) in bands.items():
		ws.merged_cells.add(
			CellRange(min_row=band_row, min_col=start + 1, max_row=band_row, max_col=start + span).coord
		)
	if bands:
		for i in range(len(columns)):
			if i not in covered:
				ws.merged_cells.add(
					CellRange(min_row=band_row, min_col=i + 1, max_row=header_row, max_col=i + 1).coord
				)

	for value, font in top:
		ws.append([_cell(ws, value, font=font)])
	ws.append([])

	def header_cell(value):
		return _cell(
			ws,
			value,
			font=bold,
			fill=header_fill,
			border=header_border,
			alignment=Alignment(horizontal="center", vertical="center", wrap_text=True),
		)

	if bands:
		# Band labels (Opening / Movement / Closing) above their columns; a column without a
		# band has its header here, spanning both header rows.
		cells = []
		for i, col in enumerate(columns):
			if i in bands:
				cells.append(
					_cell(
						ws,
						_(bands[i][1]),
						font=bold,
						fill=header_fill,
						border=header_border,
						alignment=Alignment(horizontal="center"),
					)
				)
			elif i in covered:
				cells.append(_cell(ws, None, fill=header_fill, border=header_border))
			else:
				cells.append(header_cell(_(col["label"])))
		ws.append(cells)
		ws.append(
			[
				header_cell(_(col["label"])) if i in covered else _cell(ws, None, border=header_border)
				for i, col in enumerate(columns)
			]
		)
	else:
		ws.append([header_cell(_(col["label"])) for col in columns])

	# The column that names a row on screen: group titles, subtotal titles and the indent.
	label_index = next((i for i, c in enumerate(columns) if c["key"] == result["label_key"]), 0)
	by_key = {c["key"]: c for c in nr["columns"]}

	include_subtotals = layout.get("include_subtotals", True)
	row = first_row
	for out in result["rows"]:
		kind = out["kind"]
		if kind == "foot" and not include_subtotals:
			continue
		values = {} if kind == "group" and not include_subtotals else out["values"]
		title = None
		if kind == "group":
			title = f"{_group_label(out, by_key)} ({out['count']})"
		elif kind == "foot":
			title = _("Total {0}").format(_group_label(out, by_key))
		cells = []
		for i, col in enumerate(columns):
			if title is not None and i == label_index:
				value = title
			else:
				value = _cell_value(values.get(col["key"]), col)
			number_format = None
			if col["key"] in number_formats and isinstance(value, int | float):
				number_format = number_formats[col["key"]]
			elif col["fieldtype"] == "Date" and isinstance(value, datetime.date):
				number_format = date_fmt
			elif col["fieldtype"] == "Datetime" and isinstance(value, datetime.datetime):
				number_format = date_fmt + " hh:mm:ss"
			alignment = None
			if i == label_index and out["level"]:
				alignment = Alignment(indent=min(out["level"], 15))
			font = bold if out["bold"] else None
			if (
				number_format is None
				and alignment is None
				and font is None
				and not (isinstance(value, str) and value.startswith(FORMULA_START))
				and not isinstance(value, datetime.date)
			):
				cells.append(value)
			else:
				cells.append(_cell(ws, value, font=font, alignment=alignment, number_format=number_format))
		# A row's outline is set before the row is written.
		level = min(out["level"], 7)
		if level or out["hidden"] or out["open"] is False:
			dimension = ws.row_dimensions[row]
			if level:
				dimension.outline_level = level
			if out["hidden"]:
				dimension.hidden = True
			if out["open"] is False:
				# A folded group or parent: its rows are in the sheet, hidden under it.
				dimension.collapsed = True
		ws.append(cells)
		row += 1

	# Grand total (ADR-008), always a wide bold row: the report's own reference row when it
	# is the total on screen, else the engine's total.
	total_fill = PatternFill("solid", fgColor="E7E6E6")
	cells = []
	for i, col in enumerate(columns):
		value = _cell_value(result["totals"].get(col["key"]), col) if col["key"] in result["totals"] else None
		if i == label_index and value is None:
			value = _("Total")
		number_format = (
			number_formats[col["key"]]
			if col["key"] in number_formats and isinstance(value, int | float)
			else None
		)
		cells.append(
			_cell(ws, value, font=bold, fill=total_fill, border=header_border, number_format=number_format)
		)
	ws.append(cells)

	buffer = BytesIO()
	wb.save(buffer)
	return buffer.getvalue()
