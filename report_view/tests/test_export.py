import json
import os
import time
from io import BytesIO
from unittest.mock import patch

import frappe
from openpyxl import load_workbook

from report_view import api
from report_view.core import engine
from report_view.core import view_state as vs
from report_view.core.errors import (
	RVExportMismatch,
	RVInvalidInput,
	RVInvalidViewState,
	RVLimitExceeded,
	RVNoExport,
)
from report_view.core.export import model as export_model
from report_view.core.export import xlsx
from report_view.core.normalize import normalize
from report_view.core.settings import empty_settings, load_settings
from report_view.tests.utils import AR_FILTERS, TB_FILTERS, TB_JS_FILTERS, RVTestCase, load_fixtures


def tb_model(**overrides) -> dict:
	model = {
		"version": 1,
		"format": "xlsx",
		"filters": TB_FILTERS,
		"js_filters": TB_JS_FILTERS,
		"prepared_report_name": None,
		"view_state": {"v": 1, "mode": "grid", "tree": {"depth": 3}},
		"scope": "all",
		"selection": None,
		"layout": {"title": "Trial Balance", "include_filters": True},
		"locale": {"lang": "ar", "rtl": True},
	}
	model.update(overrides)
	return model


def fixture_report(name: str) -> tuple[dict, dict]:
	"""NormalizedReport of a shared fixture (a real run on the development site) and its report filters."""
	case = next(c for c in load_fixtures() if c["name"] == name)
	i = case["input"]
	return normalize(i["raw"], i["filters"], i["report_settings"], i["register"]), i["filters"]


def fixture_workbook(name: str, state: dict, report_name: str = "General Ledger") -> tuple:
	"""The workbook the export builds for a fixture and a view state: (sheet, result, nr)."""
	nr, filters = fixture_report(name)
	model = export_model.parse(
		{"version": 1, "format": "xlsx", "filters": filters, "view_state": {"v": 1, **state}}, "xlsx"
	)
	export_model.check_against_report(model, nr)
	result = engine.apply(nr, model["view_state"], lang=frappe.local.lang)
	content = xlsx.build(report_name, nr, result, model, empty_settings(report_name))
	return load_workbook(BytesIO(content)).active, result, nr


def body_rows(ws) -> list[tuple]:
	"""(row number, values) of the sheet rows below the frozen header."""
	first = ws[ws.freeze_panes].row
	return [(r, [c.value for c in ws[r]]) for r in range(first, ws.max_row + 1)]


class TestExportXlsx(RVTestCase):
	def setUp(self):
		super().setUp()
		frappe.response.pop("filecontent", None)

	def _export(self, report: str, model: dict):
		api.export_xlsx(report, json.dumps(model))
		self.assertEqual(frappe.response["type"], "binary")
		return load_workbook(BytesIO(frappe.response["filecontent"]))

	def _reference_totals(self) -> dict:
		with self.as_user("accountant"):
			raw = api.run("Trial Balance", TB_FILTERS, TB_JS_FILTERS)["raw"]
		total = next(r for r in raw["result"] if r.get("account") == "'Total'")
		return {k: total[k] for k in ("debit", "credit", "closing_debit", "closing_credit")}

	def test_allowed_user_rtl_workbook(self):
		expected = {"totals": self._reference_totals()}
		old_lang = frappe.local.lang
		try:
			with self.as_user("accountant"):
				frappe.local.lang = "ar"
				wb = self._export("Trial Balance", tb_model(expected=expected))
		finally:
			frappe.local.lang = old_lang
		ws = wb.active
		self.assertTrue(ws.sheet_view.rightToLeft)
		self.assertEqual(ws["A1"].value, TB_FILTERS["company"])
		self.assertIsNotNone(ws.freeze_panes)
		# Band row (Opening / Movement / Closing) above the column headers.
		merged = {str(r) for r in ws.merged_cells.ranges}
		self.assertTrue(any(r.startswith("E") and ":F" in r for r in merged))
		# Outline levels for the tree and a bold grand total row with numbers.
		self.assertTrue(any((ws.row_dimensions[r].outline_level or 0) > 0 for r in range(1, ws.max_row + 1)))
		last = [c.value for c in ws[ws.max_row]]
		self.assertIn(expected["totals"]["debit"], last)
		self.assertTrue(ws.cell(row=ws.max_row, column=7).font.bold)
		number_cells = [c for c in ws[ws.max_row] if isinstance(c.value, int | float)]
		self.assertTrue(all(c.number_format.startswith("#,##0") for c in number_cells))

	def test_ltr_workbook_in_english(self):
		old_lang = frappe.local.lang
		try:
			with self.as_user("accountant"):
				frappe.local.lang = "en"
				wb = self._export("Trial Balance", tb_model())
		finally:
			frappe.local.lang = old_lang
		self.assertFalse(wb.active.sheet_view.rightToLeft)

	def test_mismatch_rejected(self):
		totals = self._reference_totals()
		totals["debit"] += 1
		# log_error may be deferred (redis) and would outlive the test rollback.
		with (
			patch("report_view.core.export.model.frappe.log_error") as log_error,
			self.as_user("accountant"),
			self.assertRaises(RVExportMismatch),
		):
			self._export("Trial Balance", tb_model(expected={"totals": totals}))
		log_error.assert_called_once()

	def test_invalid_view_state(self):
		with self.as_user("accountant"):
			for state in (
				{"v": 1, "sort": [{"key": "no_such_column", "dir": "asc"}]},
				{"v": 1, "unknown": 1},
				{"v": 1, "filters": [{"key": "debit", "type": "regex", "value": "x"}]},
				{"v": 2},
				{"v": 1, "mode": "pivot"},
			):
				with self.subTest(state=state), self.assertRaises(RVInvalidViewState):
					self._export("Trial Balance", tb_model(view_state=state))

	def test_invalid_model(self):
		with self.as_user("accountant"):
			for model in (tb_model(version=2), tb_model(format="pdf"), dict(tb_model(), extra=1)):
				with self.subTest(model=model), self.assertRaises(RVInvalidInput):
					self._export("Trial Balance", model)
			with self.assertRaises(RVInvalidInput):
				api.export_xlsx("Trial Balance", "x" * (257 * 1024))

	def test_group_export(self):
		model = tb_model(view_state={"v": 1, "group": {"by": ["currency"], "subtotals": False}})
		with self.as_user("accountant"):
			ws = self._export("Trial Balance", model).active
		rows = body_rows(ws)
		# The group row carries its numbers (no subtotal row), and they are the grand total.
		self.assertRegex(rows[0][1][0], r"^.+ \(\d+\)$")
		numbers = [v for v in rows[0][1] if isinstance(v, int | float)]
		self.assertEqual(numbers, [v for v in rows[-1][1] if isinstance(v, int | float)])
		self.assertTrue(all(ws.row_dimensions[r].outline_level == 1 for r, _values in rows[1:-1]))

	def test_expected_is_what_the_screen_shows(self):
		"""§4.2 (7): row_count = facts that pass the filters; totals = the grand total shown."""
		with self.as_user("accountant"):
			raw = api.run("Trial Balance", TB_FILTERS, TB_JS_FILTERS)["raw"]
		leaves = [r for r in raw["result"] if r.get("account") != "'Total'"]
		facts = sum(
			1 for i, r in enumerate(leaves) if i + 1 == len(leaves) or leaves[i + 1]["indent"] <= r["indent"]
		)
		expected = {"row_count": facts, "totals": self._reference_totals()}
		with self.as_user("accountant"):
			self._export("Trial Balance", tb_model(expected=expected))
			with (
				patch("report_view.core.export.model.frappe.log_error"),
				self.assertRaises(RVExportMismatch),
			):
				self._export("Trial Balance", tb_model(expected={**expected, "row_count": facts + 1}))

	def test_tree_fold_is_written_as_outline(self):
		"""tree.depth: every row is in the sheet; rows under a folded parent are hidden."""
		with self.as_user("accountant"):
			full = self._export("Trial Balance", tb_model(view_state={"v": 1, "tree": {"depth": 9}})).active
			ws = self._export("Trial Balance", tb_model(view_state={"v": 1, "tree": {"depth": 1}})).active
		rows = body_rows(ws)[:-1]
		dims = [ws.row_dimensions[r] for r, _values in rows]
		levels = [d.outline_level or 0 for d in dims]
		self.assertGreater(max(levels), 1)
		# Depth 1: the roots and their children are shown, everything deeper is hidden.
		self.assertTrue(all(bool(d.hidden) == (level > 1) for d, level in zip(dims, levels, strict=True)))
		self.assertTrue(any(d.collapsed for d in dims))
		# A folded parent is marked collapsed: at depth 1 that is every parent below the roots.
		shown_folded = [level for d, level in zip(dims, levels, strict=True) if d.collapsed and not d.hidden]
		self.assertTrue(shown_folded)
		self.assertEqual(set(shown_folded), {1})
		# Folding hides rows, it removes none: the open roots add their subtotal rows only.
		labels = [values[0] for _r, values in rows]
		feet = [label for label in labels if isinstance(label, str) and label.startswith("Total ")]
		self.assertTrue(feet)
		self.assertEqual(len(body_rows(full)[:-1]) - len(rows), self._open_parents(full) - len(feet))
		# The grand total row is the report's own 'Total' row.
		last = body_rows(ws)[-1][1]
		self.assertEqual(last[0], "Total")
		self.assertIn(self._reference_totals()["closing_debit"], last)

	def _open_parents(self, ws) -> int:
		labels = [values[0] for _r, values in body_rows(ws)[:-1]]
		return sum(1 for label in labels if isinstance(label, str) and label.startswith("Total "))

	def test_group_fold_and_subtotals(self):
		state = {"v": 1, "group": {"by": ["currency"], "subtotals": True, "expanded": "none"}}
		with self.as_user("accountant"):
			folded = self._export("Trial Balance", tb_model(view_state=state)).active
			state["group"]["expanded"] = "all"
			opened = self._export("Trial Balance", tb_model(view_state=state)).active
		rows = body_rows(folded)[:-1]
		head = folded.row_dimensions[rows[0][0]]
		self.assertTrue(head.collapsed and not head.hidden)
		self.assertTrue(all(folded.row_dimensions[r].hidden for r, _values in rows[1:]))
		# A folded group carries its numbers; an open one leaves them to its subtotal row.
		self.assertTrue(any(isinstance(v, int | float) for v in rows[0][1][1:]))
		open_rows = body_rows(opened)[:-1]
		self.assertFalse(any(isinstance(v, int | float) for v in open_rows[0][1][1:]))
		self.assertTrue(str(open_rows[-1][1][0]).startswith("Total "))
		self.assertEqual(
			[v for v in open_rows[-1][1] if isinstance(v, int | float)],
			[v for v in rows[0][1] if isinstance(v, int | float)],
		)
		self.assertEqual(len(open_rows), len(rows) + 1)

	def test_filters_the_screen_does_not_apply_are_refused(self):
		"""§4.2 (4): refused with a reason instead of a file that differs from the screen."""
		filters = (
			{"key": "account", "type": "text", "op": "not_contains", "value": "x"},
			{"key": "*", "type": "text", "op": "contains", "value": "x", "source": "search"},
			{"key": "debit", "type": "text", "op": "contains", "value": "1"},
			{"key": "debit", "type": "set", "values": [1], "exclude": False},
			{"key": "account", "type": "range", "min": 1, "max": None},
			{"key": "account", "type": "date", "from": "2026-01-01", "to": None},
			{"key": "account", "type": "set", "values": [None, "x"], "exclude": False},
		)
		with self.as_user("accountant"):
			for flt in filters:
				with self.subTest(filter=flt), self.assertRaises(RVInvalidViewState):
					self._export("Trial Balance", tb_model(view_state={"v": 1, "filters": [flt]}))
		# The same state is a valid view state (a saved view may hold it); only the export refuses.
		vs.validate({"v": 1, "filters": list(filters[:2])})
		nr, _filters = fixture_report("general_ledger_consolidated")
		for state in (
			{"v": 1, "group": {"by": ["posting_date__month"]}},
			{"v": 1, "sort": [{"key": "posting_date__year", "dir": "asc"}]},
		):
			vs.check_against_columns(state, nr["columns"])
			with self.subTest(state=state), self.assertRaises(RVInvalidViewState):
				vs.check_exportable(state, nr["columns"])
		vs.check_exportable(
			{"v": 1, "filters": [{"key": "account", "type": "set", "values": [None, "x"], "exclude": True}]},
			nr["columns"],
		)

	def test_fold_lists_are_validated(self):
		for state in (
			{"v": 1, "group": {"by": [], "expanded": [["a", 1, None, True]]}},
			{"v": 1, "tree": {"depth": None, "expanded": [0, 5]}},
		):
			vs.validate(state)
		for state in (
			{"v": 1, "group": {"by": [], "expanded": ["a"]}},
			{"v": 1, "group": {"by": [], "expanded": [[{"a": 1}]]}},
			{"v": 1, "group": {"by": [], "expanded": [[]]}},
			{"v": 1, "tree": {"expanded": ["0"]}},
			{"v": 1, "tree": {"expanded": [-1]}},
			{"v": 1, "tree": {"expanded": [1.5]}},
		):
			with self.subTest(state=state), self.assertRaises(RVInvalidViewState):
				vs.validate(state)

	def test_general_ledger_rows_and_total(self):
		"""Opening first, closing last, and the report's own 'Total' row as the grand total."""
		with self.as_user("accountant"):
			for sort in ([], [{"key": "debit", "dir": "desc"}]):
				ws, result, nr = fixture_workbook("general_ledger_consolidated", {"sort": sort})
				rows = body_rows(ws)
				keys = [c["key"] for c in result["columns"]]
				account, balance, debit = keys.index("account"), keys.index("balance"), keys.index("debit")
				reference = nr["report_totals"]["internal"][-1]
				with self.subTest(sort=sort):
					self.assertEqual(len(rows), len(nr["facts"]) + 3)
					self.assertEqual(rows[0][1][account], "Opening")
					self.assertEqual(rows[-2][1][account], "Closing (Opening + Total)")
					self.assertEqual(rows[-1][1][balance], reference["balance"])
					self.assertEqual(rows[-1][1][debit], reference["debit"])
			debits = [values[debit] or 0 for _r, values in rows[1:-2]]
			self.assertEqual(debits, sorted(debits, reverse=True))
			# With a filter the report's own rows leave and the total is the engine's.
			flt = {"key": "debit", "type": "range", "min": 1, "max": None}
			ws, result, nr = fixture_workbook("general_ledger_consolidated", {"filters": [flt]})
			rows = body_rows(ws)
			self.assertEqual(len(rows), result["row_count"] + 1)
			self.assertAlmostEqual(rows[-1][1][debit], sum(v[debit] for _r, v in rows[:-1]), places=2)

	def test_pinned_columns_and_group_rows_in_the_sheet(self):
		parties = None
		with self.as_user("accountant"):
			nr, _filters = fixture_report("accounts_receivable_group_by_party")
			parties = list(dict.fromkeys(r["party"] for r in nr["facts"]))
			state = {
				"columns": {"pinned": {"start": ["voucher_no", "party"]}},
				"group": {"by": ["party"], "subtotals": True, "expanded": [[parties[0]]]},
			}
			ws, result, nr = fixture_workbook(
				"accounts_receivable_group_by_party", state, "Accounts Receivable"
			)
		header = [c.value for c in ws[ws[ws.freeze_panes].row - 1]]
		labels = {c["key"]: c["label"] for c in nr["columns"]}
		first = next(c["key"] for c in nr["columns"] if not c["hidden"])
		self.assertEqual(header[:3], [labels[first], labels["voucher_no"], labels["party"]])
		rows = body_rows(ws)[:-1]
		self.assertEqual(len(rows), len(result["rows"]))
		for (r, values), out in zip(rows, result["rows"], strict=True):
			dimension = ws.row_dimensions[r]
			self.assertEqual(bool(dimension.hidden), out["hidden"])
			self.assertEqual(bool(dimension.collapsed), out["open"] is False)
			self.assertEqual(dimension.outline_level or 0, out["level"])
			if out["kind"] == "group":
				self.assertEqual(values[0], f"{out['group_value']} ({out['count']})")
			elif out["kind"] == "foot":
				self.assertEqual(values[0], f"Total {parties[0]}")
		self.assertEqual([out["kind"] for out in result["rows"]].count("foot"), 1)

	# ------------------------------------------------- quick filters (ADR-013)

	DEBIT_CLOSING = 'abs("closing_debit") > 0.005'

	def _leaves(self, raw: dict) -> list[dict]:
		rows = [r for r in raw["result"] if r.get("account") != "'Total'"]
		return [r for i, r in enumerate(rows) if i + 1 == len(rows) or rows[i + 1]["indent"] <= r["indent"]]

	def test_quick_filter_export_equals_the_filtered_totals(self):
		with self.as_user("accountant"):
			raw = api.run("Trial Balance", TB_FILTERS, TB_JS_FILTERS)["raw"]
		leaves = self._leaves(raw)
		kept = [r for r in leaves if abs(r.get("closing_debit") or 0) > 0.005]
		self.assertTrue(0 < len(kept) < len(leaves))
		keys = ("debit", "credit", "closing_debit", "closing_credit")
		expected = {"row_count": len(kept), "totals": {k: sum(r.get(k) or 0 for r in kept) for k in keys}}
		model = tb_model(
			view_state={"v": 1, "quick_filters": ["debit_closing"]},
			quick_filters=[
				{"id": "debit_closing", "expression": self.DEBIT_CLOSING},
				{"id": "unused", "expression": '"debit" > 1'},
			],
			expected=expected,
		)
		with self.as_user("accountant"):
			ws = self._export("Trial Balance", model).active
		rows = body_rows(ws)
		# The leaves of the file are the kept ones; the grand total is their sum.
		nr = normalize(raw, TB_FILTERS, load_settings("Trial Balance")["report_settings"])
		keys = [c["key"] for c in nr["columns"] if not c["hidden"] or c["export_hidden"]]
		account, closing_debit = keys.index("account"), keys.index("closing_debit")
		written = [values[account] for _r, values in rows[:-1]]
		self.assertTrue(all(r["account"] in written for r in kept))
		dropped = [r["account"] for r in leaves if r not in kept]
		self.assertTrue(dropped and not any(a in written for a in dropped))
		self.assertAlmostEqual(rows[-1][1][closing_debit], expected["totals"]["closing_debit"], places=2)
		# What the screen would send for the unfiltered tree no longer matches: refused, not written.
		with (
			patch("report_view.core.export.model.frappe.log_error"),
			self.as_user("accountant"),
			self.assertRaises(RVExportMismatch),
		):
			self._export("Trial Balance", dict(model, expected={"totals": self._reference_totals()}))

	def test_quick_filter_definitions_are_validated(self):
		cases = (
			# an active id without a definition
			({"v": 1, "quick_filters": ["debit_closing"]}, [], RVInvalidViewState),
			(
				{"v": 1, "quick_filters": ["q"]},
				[{"id": "x", "expression": self.DEBIT_CLOSING}],
				RVInvalidViewState,
			),
			# outside the grammar
			(
				{"v": 1, "quick_filters": ["q"]},
				[{"id": "q", "expression": 'not("debit" > 0)'}],
				RVInvalidViewState,
			),
			# unknown column, text column
			({"v": 1, "quick_filters": ["q"]}, [{"id": "q", "expression": '"nope" > 0'}], RVInvalidViewState),
			(
				{"v": 1, "quick_filters": ["q"]},
				[{"id": "q", "expression": 'abs("account") > 0'}],
				RVInvalidViewState,
			),
			# shape
			({"v": 1}, {"id": "q"}, RVInvalidInput),
			({"v": 1}, [{"id": "q"}], RVInvalidInput),
			({"v": 1}, [{"id": "q", "expression": '"debit" > 0', "label": "x"}], RVInvalidInput),
			({"v": 1}, [{"id": "q", "expression": 1}], RVInvalidInput),
			({"v": 1}, [{"id": "q", "expression": '"debit" > 0'}] * 2, RVInvalidInput),
			({"v": 1}, [{"id": str(i), "expression": '"debit" > 0'} for i in range(51)], RVInvalidInput),
		)
		with self.as_user("accountant"):
			for state, definitions, error in cases:
				with self.subTest(state=state, definitions=definitions), self.assertRaises(error):
					self._export("Trial Balance", tb_model(view_state=state, quick_filters=definitions))
			# An inactive definition that refers to a column this report lacks is not checked:
			# the browser does not apply it either.
			self._export(
				"Trial Balance",
				tb_model(view_state={"v": 1}, quick_filters=[{"id": "q", "expression": '"nope" > 0'}]),
			)

	def test_quick_filter_denied_user(self):
		model = tb_model(
			view_state={"v": 1, "quick_filters": ["debit_closing"]},
			quick_filters=[{"id": "debit_closing", "expression": self.DEBIT_CLOSING}],
		)
		with self.as_user("none"), self.assertRaises(frappe.PermissionError):
			self._export("Trial Balance", model)

	def test_user_without_export_rejected(self):
		model = tb_model(filters=AR_FILTERS, js_filters=[], view_state={"v": 1})
		with self.as_user("accountant"):
			if frappe.permissions.can_export("Sales Invoice"):
				self.skipTest("Accounts User can export Sales Invoice on this site")
			with self.assertRaises(RVNoExport):
				self._export("Accounts Receivable", model)

	def test_denied_user(self):
		with self.as_user("none"), self.assertRaises(frappe.PermissionError):
			self._export("Trial Balance", tb_model())


class TestExportSafetyAndJobs(RVTestCase):
	"""Text never becomes a formula; a large export is a background job of its user."""

	def setUp(self):
		super().setUp()
		frappe.response.pop("filecontent", None)
		self._free_users()

	def tearDown(self):
		self._free_users()
		super().tearDown()

	def _free_users(self):
		# the job state lives in Redis, outside the rolled-back transaction
		from report_view.core.export import jobs

		for user in self.users.values():
			frappe.cache.delete_value(jobs.JOBS.active_key(user))

	def test_text_that_looks_like_a_formula_stays_text(self):
		columns = [
			{"key": "party", "label": "=Party", "fieldtype": "Data", "width": 120},
			{"key": "amount", "label": "Amount", "fieldtype": "Float", "width": 120},
		]
		rows = [
			{"party": '=HYPERLINK("http://x","a")', "amount": -5},
			{"party": "+cmd|' /C calc'!A0", "amount": 1.5},
			{"party": "-2+3", "amount": None},
			{"party": "@SUM(A1)", "amount": 0},
			{"party": "\tTab", "amount": 2},
			{"party": "Plain - text", "amount": -0.25},
		]
		result = {
			"columns": columns,
			"label_key": "party",
			"rows": [
				{
					"kind": "row",
					"values": r,
					"level": 0,
					"bold": False,
					"hidden": False,
					"open": None,
					"label": None,
				}
				for r in rows
			],
			"totals": {"amount": -1.75},
		}
		model = {
			"layout": {"title": "=1+1", "include_filters": True},
			"filters": {"company": "=Co", "party": "@x"},
		}
		content = xlsx.build(
			"Trial Balance", {"columns": columns}, result, model, empty_settings("Trial Balance")
		)
		ws = load_workbook(BytesIO(content)).active
		cells = [c for row in ws.iter_rows() for c in row if c.value is not None]
		self.assertFalse([c.coordinate for c in cells if c.data_type == "f"])
		self.assertEqual(ws["A1"].value, "=Co")
		self.assertTrue(ws["A1"].quotePrefix)
		self.assertEqual(ws["A2"].value, "=1+1")
		body = body_rows(ws)
		self.assertEqual([v[0] for _r, v in body[:-1]], [r["party"] for r in rows])
		self.assertTrue(all(ws.cell(row=r, column=1).quotePrefix for r, _v in body[:5]))
		# plain text and real numbers are written as before
		self.assertFalse(ws.cell(row=body[5][0], column=1).quotePrefix)
		self.assertEqual([v[1] for _r, v in body[:-1]], [-5, 1.5, None, 0, 2, -0.25])
		self.assertEqual(ws.cell(row=body[0][0], column=2).data_type, "n")
		self.assertEqual(body[-1][1], ["Total", -1.75])

	def test_a_large_export_is_a_job_of_its_user(self):
		from report_view.core.export import jobs

		with (
			self.as_user("accountant"),
			patch.object(jobs, "ASYNC_ROWS", 5),
			patch("report_view.core.export.jobs.frappe.enqueue") as enqueue,
		):
			out = api.export_xlsx("Trial Balance", json.dumps(tb_model()))
			self.assertEqual(out["status"], "queued")
			self.assertGreater(out["row_count"], 5)
			self.assertNotIn("filecontent", frappe.response)
			enqueue.assert_called_once()
			self.assertEqual(enqueue.call_args.kwargs["queue"], "long")
			job = out["job"]
			self.assertEqual(api.export_xlsx_status(job)["state"], "queued")
			# one export job at a time per user
			with self.assertRaises(RVLimitExceeded):
				api.export_xlsx("Trial Balance", json.dumps(tb_model()))
			# the browser's expected row count queues without a run in the request
			with patch("report_view.core.export.jobs.runner.run") as run:
				with self.assertRaises(RVLimitExceeded):
					api.export_xlsx("Trial Balance", json.dumps(tb_model(expected={"row_count": 6})))
				run.assert_not_called()
			with self.assertRaises(RVInvalidInput):
				api.export_xlsx_download(job)
			old_lang = frappe.local.lang
			try:
				jobs.run_job(job)
			finally:
				frappe.local.lang = old_lang
			self.assertEqual(api.export_xlsx_status(job)["state"], "done")
		# another user's job does not exist for them
		with self.as_user("admin"), self.assertRaises(frappe.DoesNotExistError):
			api.export_xlsx_status(job)
		with self.as_user("accountant"):
			api.export_xlsx_download(job)
			self.assertEqual(frappe.response["type"], "binary")
			self.assertEqual(frappe.response["filename"], "trial_balance.xlsx")
			ws = load_workbook(BytesIO(frappe.response["filecontent"])).active
			self.assertEqual(ws["A1"].value, TB_FILTERS["company"])
			# downloaded once: the job and its file are gone
			self.assertFalse(os.path.exists(jobs._path(job)))
			with self.assertRaises(frappe.DoesNotExistError):
				api.export_xlsx_status(job)
			for bad in ("x", "../../etc", "A" * 16):
				with self.assertRaises(RVInvalidInput):
					api.export_xlsx_status(bad)

	def test_a_failed_job_says_why_and_frees_the_user(self):
		from report_view.core.export import jobs

		with (
			self.as_user("accountant"),
			patch.object(jobs, "ASYNC_ROWS", 5),
			patch("report_view.core.export.jobs.frappe.enqueue"),
		):
			totals = {"debit": -1}
			job = api.export_xlsx("Trial Balance", json.dumps(tb_model(expected={"totals": totals})))["job"]
			with patch("report_view.core.export.model.frappe.log_error"):
				jobs.run_job(job)
			state = api.export_xlsx_status(job)
			self.assertEqual(state["state"], "failed")
			self.assertTrue(state["error"])
			# the user may start another export
			self.assertEqual(api.export_xlsx("Trial Balance", json.dumps(tb_model()))["status"], "queued")

	def test_a_job_whose_worker_is_gone_fails_and_frees_the_user(self):
		"""C16: a killed worker leaves the job ``running``; past its timeout it reads as lost."""
		from report_view.core.export import jobs

		with (
			self.as_user("accountant"),
			patch.object(jobs, "ASYNC_ROWS", 5),
			patch("report_view.core.export.jobs.frappe.enqueue") as enqueue,
		):
			job = api.export_xlsx("Trial Balance", json.dumps(tb_model()))["job"]
			self.assertEqual(enqueue.call_args.kwargs["job_id"], jobs.JOBS.rq_id(job))
			state = jobs._get(job)
			jobs._set(
				job, {**state, "state": "running", "updated_at": time.time() - jobs.JOBS.stale_seconds - 5}
			)
			with patch("report_view.core.jobs.frappe.log_error") as log_error:
				out = api.export_xlsx_status(job)
			self.assertEqual((out["state"], out["error_code"], out["cancellable"]), ("failed", "lost", False))
			self.assertTrue(out["error"])
			log_error.assert_called_once()
			self.assertEqual(api.export_xlsx("Trial Balance", json.dumps(tb_model()))["status"], "queued")

	def test_a_job_is_cancelled_by_its_user_only(self):
		from report_view.core.export import jobs

		with (
			self.as_user("accountant"),
			patch.object(jobs, "ASYNC_ROWS", 5),
			patch("report_view.core.export.jobs.frappe.enqueue"),
		):
			job = api.export_xlsx("Trial Balance", json.dumps(tb_model()))["job"]
			self.assertTrue(api.export_xlsx_status(job)["cancellable"])
		with self.as_user("admin"), self.assertRaises(frappe.DoesNotExistError):
			api.export_xlsx_cancel(job)
		with (
			self.as_user("accountant"),
			patch.object(jobs, "ASYNC_ROWS", 5),
			patch("report_view.core.export.jobs.frappe.enqueue"),
		):
			out = api.export_xlsx_cancel(job)
			self.assertEqual(
				(out["state"], out["error_code"], out["cancellable"]), ("failed", "cancelled", False)
			)
			# a cancelled job that a worker picks up never runs, and keeps no file
			with patch("report_view.core.export.jobs.runner.run") as run:
				jobs.run_job(job)
				run.assert_not_called()
			self.assertFalse(os.path.exists(jobs._path(job)))
			self.assertEqual(api.export_xlsx_status(job)["error_code"], "cancelled")
			# the user may start another export
			self.assertEqual(api.export_xlsx("Trial Balance", json.dumps(tb_model()))["status"], "queued")

	def test_the_queue_comes_from_the_site_config(self):
		from report_view.core.export import jobs

		with patch.dict(frappe.conf, {"report_view_export_queue": "rv_xlsx_q"}):
			self.assertEqual(jobs.JOBS.queue(), "long")
			with patch(
				"frappe.utils.background_jobs.get_queues_timeout",
				return_value={"short": 300, "default": 300, "long": 1500, "rv_xlsx_q": 1500},
			):
				self.assertEqual(jobs.JOBS.queue(), "rv_xlsx_q")

	def test_cleanup_removes_old_files_only(self):
		from report_view.core.export import jobs

		os.makedirs(jobs._path(None), exist_ok=True)
		old, new = jobs._path("0" * 16), jobs._path("1" * 16)
		for path in (old, new):
			with open(path, "wb") as f:
				f.write(b"x")
		os.utime(old, (0, 0))
		jobs.cleanup()
		self.assertFalse(os.path.exists(old))
		self.assertTrue(os.path.exists(new))
		os.remove(new)
